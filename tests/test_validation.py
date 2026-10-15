"""Regression tests for observable edits and template preservation. No Office UI claim."""
import contextlib
import copy
import importlib.util
import io
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from lxml import etree as ET
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "ssr-ppt-skill" / "scripts"
spec = importlib.util.spec_from_file_location("preview", SCRIPTS / "preview.py")
preview = importlib.util.module_from_spec(spec)
spec.loader.exec_module(preview)


def add_chart(slide, kind=XL_CHART_TYPE.COLUMN_CLUSTERED):
    data = CategoryChartData()
    data.categories = ["第一组", "第二组", "第三组"]
    data.add_series("已完成", [6, 4, 3])
    return slide.shapes.add_chart(kind, Inches(1), Inches(2), Inches(6), Inches(3), data)


def multi_master_fixture(path):
    """Two linked masters/themes, notes and hyperlinks, generated without customer data."""
    prs = Presentation()
    for name in ("保留标题", "第二主题"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = name
        run.hyperlink.address = "https://example.com/fixture"
        slide.notes_slide.notes_text_frame.text = "备注：固定口径与来源"
    prs.save(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
    def root(name):
        return ET.fromstring(parts[name])
    def save(name, element):
        parts[name] = ET.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)
    # Clone the blank layout and its master, keeping a single valid layout relationship.
    master = root("ppt/slideMasters/slideMaster1.xml")
    rels = root("ppt/slideMasters/_rels/slideMaster1.xml.rels")
    keep = next(r.get("Id") for r in rels if r.get("Target") == "../slideLayouts/slideLayout7.xml")
    for rel in list(rels):
        if rel.get("Type").endswith("/slideLayout"):
            if rel.get("Id") == keep:
                rel.set("Target", "../slideLayouts/slideLayout12.xml")
            else:
                rels.remove(rel)
        elif rel.get("Type").endswith("/theme"):
            rel.set("Target", "../theme/theme2.xml")
    layouts = master.find(f"{{{pns}}}sldLayoutIdLst")
    for item in list(layouts):
        if item.get(f"{{{rns}}}id") != keep:
            layouts.remove(item)
    save("ppt/slideMasters/slideMaster2.xml", master)
    save("ppt/slideMasters/_rels/slideMaster2.xml.rels", rels)
    parts["ppt/theme/theme2.xml"] = parts["ppt/theme/theme1.xml"].replace(b'Calibri', b'Arial')
    parts["ppt/slideLayouts/slideLayout12.xml"] = parts["ppt/slideLayouts/slideLayout7.xml"]
    parts["ppt/slideLayouts/_rels/slideLayout12.xml.rels"] = parts["ppt/slideLayouts/_rels/slideLayout7.xml.rels"].replace(b'slideMaster1.xml', b'slideMaster2.xml')
    parts["ppt/slides/_rels/slide2.xml.rels"] = parts["ppt/slides/_rels/slide2.xml.rels"].replace(b'slideLayout7.xml', b'slideLayout12.xml')
    rels = root("ppt/_rels/presentation.xml.rels")
    ET.SubElement(rels, f"{{{relns}}}Relationship", Id="rIdTemplate2", Type=rns+"/slideMaster", Target="slideMasters/slideMaster2.xml")
    save("ppt/_rels/presentation.xml.rels", rels)
    presentation = root("ppt/presentation.xml")
    item = ET.SubElement(presentation.find(f"{{{pns}}}sldMasterIdLst"), f"{{{pns}}}sldMasterId", id="2147483649")
    item.set(f"{{{rns}}}id", "rIdTemplate2")
    save("ppt/presentation.xml", presentation)
    types = root("[Content_Types].xml")
    for part, suffix in (("slideMasters/slideMaster2.xml", "presentationml.slideMaster+xml"), ("slideLayouts/slideLayout12.xml", "presentationml.slideLayout+xml"), ("theme/theme2.xml", "theme+xml")):
        ET.SubElement(types, f"{{{ctns}}}Override", PartName="/ppt/"+part, ContentType="application/vnd.openxmlformats-officedocument."+suffix)
    save("[Content_Types].xml", types)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


class EditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.work = Path(self.tmp.name)
        self.src = self.work / "source.pptx"

    def fixture(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = "页面标题"
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.text = "组内关键数值 12"
        node = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(1), Inches(1))
        line = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(1), Inches(4), Inches(1))
        line.begin_connect(box, 3)
        line.end_connect(node, 1)
        table = slide.shapes.add_table(2, 2, 0, Inches(3), Inches(3), Inches(2))
        table.table.cell(0, 0).text = "原先已有 ✓"
        image = self.work / "picture.png"
        Image.new("RGB", (100, 60), "#256360").save(image)
        slide.shapes.add_picture(str(image), Inches(6), Inches(1), width=Inches(2))
        prs.save(self.src)
        return group.shape_id, box.shape_id, table.shape_id

    def test_group_table_picture_roundtrip_and_source_unchanged(self):
        gid, tid, tableid = self.fixture()
        original = self.src.read_bytes()
        result = preview.edit_test(self.src, self.work / "out")
        self.assertTrue(result["checks"])
        self.assertTrue(all(c["status"] == "passed" for c in result["checks"]))
        self.assertTrue(any(c["target"]["path"] == [gid, tid] for c in result["checks"]))
        self.assertEqual(self.src.read_bytes(), original)
        # Existing checkmarks and unchanged group text must not mask a failed edit.
        reloaded = Presentation(result["copy"])
        for c in result["checks"]:
            if c["target"]["kind"] == "text" and c["target"]["path"] == [gid, tid]:
                preview.locate_shape(reloaded, c["target"]).text = c["before"]
            if c["target"]["path"] == [tableid]:
                preview.locate_shape(reloaded, c["target"]).table.cell(0, 0).text = c["before"]
        checked = preview.verify_edits(reloaded, copy.deepcopy(result["checks"]))
        self.assertEqual(sum(c["status"] == "failed" for c in checked), 2)

    def test_supported_chart_values_and_workbook(self):
        for kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE):
            with self.subTest(kind=kind):
                prs = Presentation()
                shape = add_chart(prs.slides.add_slide(prs.slide_layouts[6]), kind)
                prs.save(self.src)
                target = {"slide": 1, "path": [shape.shape_id], "kind": "chart"}
                result = preview.edit_test(self.src, self.work / str(kind.value), targets=[target])
                record = result["checks"][0]
                self.assertEqual(record["status"], "passed", record)
                self.assertEqual(record["actual"]["series"][0]["values"][0], 7.6)
                self.assertEqual(record["actual"]["workbook"]["B2"], 7.6)
                unchanged = preview.verify_edits(Presentation(self.src), copy.deepcopy(result["checks"]))
                self.assertEqual(unchanged[0]["status"], "failed")
        # A cached number can hide a formula; do not silently replace its workbook.
        with zipfile.ZipFile(BytesIO(shape.chart.part.chart_workbook.xlsx_part.blob)) as source:
            parts = {n: source.read(n) for n in source.namelist()}
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        sheet = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = sheet.find(f".//{{{ns}}}c[@r='B2']")
        ET.SubElement(cell, f"{{{ns}}}f").text = "2+4"
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
        blob = BytesIO()
        with zipfile.ZipFile(blob, "w") as output:
            for name, data in parts.items():
                output.writestr(name, data)
        with self.assertRaises(NotImplementedError):
            preview.workbook_blob_cells(blob.getvalue())

    def test_empty_workbook_value_cell_does_not_crash(self):
        # PptxGenJS writes empty <v/> cells in embedded workbooks; they must read as empty, not crash.
        prs = Presentation()
        shape = add_chart(prs.slides.add_slide(prs.slide_layouts[6]))
        with zipfile.ZipFile(BytesIO(shape.chart.part.chart_workbook.xlsx_part.blob)) as source:
            parts = {n: source.read(n) for n in source.namelist()}
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        sheet = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        row = sheet.find(f".//{{{ns}}}row")
        ET.SubElement(ET.SubElement(row, f"{{{ns}}}c", r="Z1"), f"{{{ns}}}v")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
        blob = BytesIO()
        with zipfile.ZipFile(blob, "w") as output:
            for name, data in parts.items():
                output.writestr(name, data)
        cells = preview.workbook_blob_cells(blob.getvalue())
        self.assertNotIn("Z1", cells)
        self.assertEqual(cells["B2"], 6)

    def test_unsupported_chart_is_not_passed(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = XyChartData()
        data.add_series("XY").add_data_point(1, 2)
        slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, 0, 0, Inches(5), Inches(3), data)
        prs.save(self.src)
        result = preview.edit_test(self.src, self.work / "out")
        self.assertEqual(result["checks"][0]["status"], "unsupported")
        self.assertEqual(preview.checks_status(result["checks"]), "unsupported")

    def test_explicit_target_and_invalid_selector(self):
        gid, tid, _ = self.fixture()
        target = {"slide": 1, "path": [gid, tid], "kind": "text"}
        result = preview.edit_test(self.src, self.work / "out", targets=[target])
        self.assertEqual(len(result["checks"]), 1)
        self.assertEqual(result["checks"][0]["status"], "passed")
        target["slide"] = 0
        result = preview.edit_test(self.src, self.work / "bad", targets=[target])
        self.assertEqual(result["checks"][0]["status"], "failed")

    def test_strict_statuses(self):
        for status in ("failed", "not_run", "unsupported"):
            with self.subTest(status=status), self.assertRaises(SystemExit):
                preview.write_report({"checks": [preview.check("required", status)]}, self.work, True)
        report = {"checks": [preview.check("required", "passed"), preview.check("unused", "not_applicable", required=False)]}
        preview.write_report(report, self.work, True)
        self.assertEqual(report["automatic_status"], "passed")
        self.assertEqual(report["manual_review"]["status"], "not_run")

    def test_multi_theme_inventory_and_scoped_font_change(self):
        multi_master_fixture(self.src)
        original = self.src.read_bytes()
        themes = preview.font_inventory(self.src)["themes"]
        self.assertEqual(len(themes), 2)
        out = self.work / "font.pptx"
        args = [sys.executable, str(SCRIPTS / "set_theme_fonts.py"), str(self.src), "--heading", "Test Heading", "--body", "Test Body", "-o", str(out)]
        result = subprocess.run(args, capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(out.exists())
        result = subprocess.run(args+["--theme", "ppt/theme/theme1.xml"], capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(out) as b:
            self.assertNotEqual(a.read("ppt/theme/theme1.xml"), b.read("ppt/theme/theme1.xml"))
            for name in a.namelist():
                if name != "ppt/theme/theme1.xml":
                    self.assertEqual(a.read(name), b.read(name), name)
        self.assertEqual(self.src.read_bytes(), original)

    def test_template_roundtrip_preserves_other_parts(self):
        multi_master_fixture(self.src)
        prs = Presentation(self.src)
        self.assertEqual(len(prs.slide_masters), 2)
        self.assertEqual(prs.slides[1].slide_layout.part.partname, "/ppt/slideLayouts/slideLayout12.xml")
        prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "只改这句话"
        out = self.work / "roundtrip.pptx"
        prs.save(out)
        reloaded = Presentation(out)
        self.assertEqual(reloaded.slides[0].shapes[0].text, "只改这句话")
        self.assertEqual(reloaded.slides[1].shapes[0].text, "第二主题")
        for slide in reloaded.slides:
            self.assertIn("固定口径", slide.notes_slide.notes_text_frame.text)
            self.assertEqual(slide.shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address, "https://example.com/fixture")
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(out) as b:
            for name in a.namelist():
                if name.startswith(("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/notesSlides/")) or name == "ppt/slides/slide2.xml":
                    self.assertEqual(a.read(name), b.read(name), name)


class KitTests(unittest.TestCase):
    def test_pptx_kit_defaults_and_editability(self):
        sys.path.insert(0, str(SCRIPTS))
        try:
            from pptx_kit import Deck
        finally:
            sys.path.remove(str(SCRIPTS))
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            Image.new("RGB", (1600, 600), "#256360").save(work / "wide.png")
            deck = Deck()
            s = deck.slide(bg="FBF9F4")
            deck.text(s, 80, 80, 800, 60, [[("改变", {"color": "E8590C"}), ("城市零售", {})]], size=44, font="Test Serif")
            chart = deck.bar_chart(s, 80, 200, 700, 300, ["甲", "乙", "丙"], [5, 3, 2], highlight={0}, font="Test Sans")
            table = deck.table(s, 80, 520, 600, [["城市", "数值"], ["上海", "86.4"]], align=["left", "right"])
            pic = deck.picture(s, work / "wide.png", 800, 200, 300, 300)
            box = deck.rect(s, 800, 520, 100, 50, fill="E8590C")
            src = deck.save(work / "kit.pptx", theme_fonts=("Test Serif", "Test Sans"))

            prs = Presentation(src)
            slide = prs.slides[0]
            self.assertEqual((prs.slide_width, prs.slide_height), (12192000, 6858000))
            runs = [r for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
            self.assertTrue(runs and all(r._r.find(".//" + qn_ea) is not None for r in runs))
            plot = [sh for sh in slide.shapes if sh.has_chart][0].chart.plots[0]
            self.assertEqual(len(plot.series), 1)
            self.assertEqual(list(plot.series[0].values), [5, 3, 2])  # data untouched by the highlight
            self.assertIn("2D5ABB26", ET.tostring([sh for sh in slide.shapes if sh.has_table][0]._element).decode())
            # 1600x600 into a square keeps the centre 600/1600 of the width: each side loses (1 - 0.375) / 2
            self.assertAlmostEqual(next(sh for sh in slide.shapes if sh.shape_id == pic.shape_id).crop_left, 0.3125, places=4)
            rect = next(sh for sh in slide.shapes if sh.shape_id == box.shape_id)._element
            self.assertIsNotNone(rect.find(".//" + qn_effect))
            self.assertIsNone(rect.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style"))  # no theme shadow
            with zipfile.ZipFile(src) as z:
                self.assertIn(b'<a:ea typeface="Test Sans"/>', z.read("ppt/theme/theme1.xml"))
            result = preview.edit_test(src, work / "edit", targets=[
                {"slide": 1, "path": [chart.shape_id], "kind": "chart"},
                {"slide": 1, "path": [table.shape_id], "kind": "table", "cell": [1, 1]}])
            self.assertEqual([c["status"] for c in result["checks"]], ["passed", "passed"])


    def test_bar_chart_axis_covers_negative_zero_and_mixed_values(self):
        sys.path.insert(0, str(SCRIPTS))
        try:
            from pptx_kit import Deck
        finally:
            sys.path.remove(str(SCRIPTS))
        for values in ([-10, -5], [0, 0, 0], [-3, 5], [34, 24, 15], [7], [None, 4]):
            with self.subTest(values=values):
                deck = Deck()
                axis = deck.bar_chart(deck.slide(), 80, 80, 600, 300, [str(i) for i in range(len(values))],
                                      values).chart.value_axis
                numeric = [v for v in values if v is not None]
                low, high = axis.minimum_scale, axis.maximum_scale
                self.assertLess(low, high)
                self.assertLessEqual(low, min(min(numeric), 0))
                self.assertGreaterEqual(high, max(max(numeric), 0))
        # Negative bars must not be inverted (rendered white or mirrored), for the series and highlighted points.
        deck = Deck()
        frame = deck.bar_chart(deck.slide(), 80, 80, 600, 300, ["a", "b"], [-3, 5], highlight={0})
        xml = ET.tostring(frame.chart._chartSpace).decode()
        self.assertIn('<c:dPt><c:idx val="0"/><c:invertIfNegative val="0"/>', xml)
        self.assertIn('<c:invertIfNegative val="0"/>', xml.split("<c:dPt>")[0])
        self.assertIn('<c:tickLblPos val="low"/>', xml)
        with self.assertRaises(ValueError):
            Deck().bar_chart(Deck().slide(), 0, 0, 100, 100, ["a"], [None])

    def test_icons_are_native_line_shapes_and_svg_becomes_a_picture(self):
        sys.path.insert(0, str(SCRIPTS))
        try:
            import pptx_kit
        finally:
            sys.path.remove(str(SCRIPTS))
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        # Path parsing: relative commands, H/V, packed arc flags ("012 2") and closepath.
        subs = pptx_kit.svg_path_to_subpaths("M1 1h4v4H1zm8 0a2 2 0 012 2")
        self.assertEqual(subs[0][-2:], [("L", 1.0, 5.0), ("Z",)])
        self.assertEqual(subs[1][0], ("M", 9.0, 1.0))
        self.assertEqual(subs[1][-1][0], "C")
        self.assertAlmostEqual(subs[1][-1][-2], 11.0)
        self.assertAlmostEqual(subs[1][-1][-1], 3.0)
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            deck = pptx_kit.Deck()
            s = deck.slide()
            icon = deck.icon(s, "factory", 80, 80, size=48, color="1E5AA8")
            custom = work / "dot.svg"
            custom.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
                              '<circle cx="5" cy="5" r="4" fill="#000"/><line x1="0" y1="0" x2="10" y2="10"/></svg>',
                              encoding="utf-8")
            deck.icon(s, custom, 200, 80, size=48, color="E8590C")
            with self.assertRaises(KeyError):
                deck.icon(s, "no-such-icon", 0, 0)
            svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 120"><rect x="10" y="40" width="80" '
                   'height="60" rx="6" fill="#E8F1FA" stroke="#1E5AA8" stroke-width="3"/></svg>')
            pic = deck.svg_picture(s, svg, 400, 80, 300, 180, out=work / "truck.png")
            src = deck.save(work / "icons.pptx")
            with Image.open(work / "truck.png") as png:
                self.assertEqual(png.size, (600, 360))                          # twice the box size
            slide = Presentation(src).slides[0]
            shape = next(sh for sh in slide.shapes if sh.shape_id == icon.shape_id)
            xml = ET.tostring(shape._element).decode()
            self.assertIn("custGeom", xml)
            self.assertNotIn("prstGeom", xml)
            self.assertNotIn("p:style", xml)                                   # no theme shadow
            self.assertIn('fill="none"', xml)                                   # line icon: paths not filled
            self.assertIn("cubicBezTo", xml)                                    # arcs became curves
            ln = shape._element.find(f".//{a}ln")
            self.assertEqual(ln.get("cap"), "rnd")
            self.assertEqual(ln.find(f".//{a}srgbClr").get("val"), "1E5AA8")
            filled = ET.tostring(slide.shapes[1]._element).decode()
            self.assertIn('<a:solidFill><a:srgbClr val="E8590C"/></a:solidFill>', filled)  # the filled circle
            self.assertTrue(any(sh.shape_id == pic.shape_id and sh.shape_type == 13 for sh in slide.shapes))


class NoRendererTests(unittest.TestCase):
    def test_pptx_check_runs_edit_test_without_libreoffice(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "标题"
            prs.save(work / "deck.pptx")
            base = [sys.executable, str(SCRIPTS / "preview.py"), "pptx", str(work / "deck.pptx"),
                    "--edit-test", "--soffice", str(work / "no-such-soffice.exe")]
            result = subprocess.run(base + ["--out", str(work / "review")], capture_output=True, timeout=120)
            self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
            report = json.loads((work / "review/report.json").read_text(encoding="utf-8"))
            checks = {c["name"]: c["status"] for c in report["checks"]}
            self.assertEqual(checks["render"], "not_run")
            self.assertEqual(checks["layout"], "not_run")
            self.assertTrue(any(n.startswith("edit:") and s == "passed" for n, s in checks.items()))
            strict = subprocess.run(base + ["--out", str(work / "strict"), "--strict"], capture_output=True, timeout=120)
            self.assertEqual(strict.returncode, 1)  # a required check was not run, so strict mode must not pass


class TextDensityTests(unittest.TestCase):
    LONG = "社区活动报名分散在多个街道，现场核对名单耗时较长，部分居民住得较远，建议活动中心统一安排志愿者就近登记，并按服务时长折算积分，同时规定回访时间。"
    POINT = "志愿者就近登记，按时长折算积分"

    def deck(self, path):
        def box(shapes, x, y, w, text, size=16):
            tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.6))
            run = tb.text_frame.paragraphs[0].add_run()
            run.text, run.font.size = text, Pt(size)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        heavy = prs.slides.add_slide(prs.slide_layouts[5])                 # title-only layout
        heavy.shapes.title.text = self.LONG                                # titles are not body text
        box(heavy.shapes, 0.8, 1.6, 6, self.LONG)                          # one long paragraph
        group = heavy.shapes.add_group_shape()                             # a card with three short points
        group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5), Inches(1.6), Inches(5), Inches(3))
        for i in range(3):
            box(group.shapes, 7.7, 1.8 + i * 0.8, 4.6, f"{i + 1}. {self.POINT}{'，并记录在案' * i}")
        box(heavy.shapes, 0.8, 4.4, 6, self.LONG[:30], size=40)            # headline-size text
        box(heavy.shapes, 0.8, 6.8, 12, "来源：" + self.LONG, size=10)     # footer band
        light = prs.slides.add_slide(prs.slide_layouts[6])
        box(light.shapes, 0.8, 1.0, 4, "发起站点")
        box(light.shapes, 6.0, 1.0, 4, "协助站点 · 7∶3")
        prs.save(path)

    def test_visual_first_hints_point_at_long_text_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "text.pptx"
            self.deck(path)
            text = preview.pptx_text(path, visual_first=True)
            heavy, light = text["slides"]
            kinds = sorted(f["kind"] for f in heavy["findings"])
            self.assertEqual(kinds, ["long_paragraph", "long_text_block"])
            self.assertEqual(preview.text_units(self.LONG), heavy["longest_paragraph"])  # title, headline, footer left out
            card = next(f for f in heavy["findings"] if f["kind"] == "long_text_block")
            self.assertEqual((card["rect"]["x"], card["rect"]["w"]), (540, 360))           # group coordinates mapped to points
            self.assertEqual(light["findings"], [])
            plain = preview.pptx_text(path, visual_first=False)
            self.assertEqual([s["findings"] for s in plain["slides"]], [[], []])            # numbers only without the mode
            self.assertEqual(plain["slides"][0]["units"], heavy["units"])

    def test_units_count_cjk_characters_and_latin_words(self):
        self.assertEqual(preview.text_units("营收 42.5 亿元，同比 +8.3%"), 6 + 2 * 2)
        self.assertEqual(preview.text_units("Queue time fell 50%"), 2 * 4)

    def test_cli_reports_hints_without_failing(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            self.deck(work / "text.pptx")
            cmd = [sys.executable, str(SCRIPTS / "preview.py"), "pptx", str(work / "text.pptx"), "--out",
                   str(work / "review"), "--soffice", str(work / "no-such-soffice.exe"), "--visual-first"]
            result = subprocess.run(cmd, capture_output=True, timeout=120)
            self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
            self.assertIn("visual-first hints: 2", result.stdout.decode("utf-8"))
            report = json.loads((work / "review/report.json").read_text(encoding="utf-8"))
            self.assertTrue(report["text"]["visual_first"])
            self.assertNotIn("text", {c["name"] for c in report["checks"]})                 # never a pass/fail check


def lines_pdf(path, pages=1):
    """A 960x540 pt PDF whose first page has a label crossed by a line, plus an underline, a strikethrough
    and a line hidden behind a filled label box, none of which should be reported."""
    pm = preview.pymupdf_or_none()
    doc = pm.open()
    for _ in range(pages):
        doc.new_page(width=960, height=540)
    page = doc[0]
    page.insert_text((100, 100), "Crossed label", fontname="helv", fontsize=20)
    page.draw_line((80, 93), (400, 93), color=(0, 0.5, 0), width=1.2)               # through the glyphs
    page.insert_text((100, 200), "Underlined note", fontname="helv", fontsize=20)
    page.draw_line((100, 203), (250, 203), color=(0, 0, 0), width=1)                # below the baseline
    width = pm.get_text_length("Old price", fontname="helv", fontsize=20)
    page.insert_text((100, 300), "Old price", fontname="helv", fontsize=20)
    page.draw_line((100, 293), (100 + width, 293), color=(0, 0, 0), width=1)        # strikethrough
    page.draw_line((80, 393), (400, 393), color=(0, 0, 0), width=1.2)               # painted first ...
    page.draw_rect(pm.Rect(95, 378, 260, 406), color=None, fill=(1, 1, 1))          # ... then covered
    page.insert_text((100, 400), "Label on a line", fontname="helv", fontsize=20)
    doc.save(path)
    doc.close()


@unittest.skipUnless(preview.pymupdf_or_none(), "PyMuPDF is needed to read PDFs")
class LineThroughTextTests(unittest.TestCase):
    def test_only_a_visible_line_through_glyphs_is_reported(self):
        pm = preview.pymupdf_or_none()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "lines.pdf"
            lines_pdf(path)
            doc = pm.open(path)
            findings = preview.pdf_layout(doc[0], pm)
            doc.close()
        self.assertEqual([(f["kind"], f["severity"], f["text"]) for f in findings],
                         [("line_through_text", "warning", "Crossed label")])

    def test_pptx_check_uses_supplied_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
            for text in ("第一页", "第二页"):
                prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = text
            prs.save(work / "deck.pptx")
            lines_pdf(work / "deck.pdf", pages=2)
            newer = (work / "deck.pptx").stat().st_mtime + 60
            os.utime(work / "deck.pdf", (newer, newer))
            base = [sys.executable, str(SCRIPTS / "preview.py"), "pptx", str(work / "deck.pptx"), "--edit-test",
                    "--soffice", str(work / "no-such-soffice.exe"), "--strict"]
            result = subprocess.run(base + ["--pdf", str(work / "deck.pdf"), "--out", str(work / "review")],
                                    capture_output=True, timeout=120)
            self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
            report = json.loads((work / "review/report.json").read_text(encoding="utf-8"))
            checks = {c["name"]: c["status"] for c in report["checks"]}
            self.assertEqual((checks["render"], checks["layout"]), ("passed", "passed"))  # a warning does not fail
            self.assertTrue(report["renderer"].startswith("supplied PDF"))
            self.assertEqual([f["kind"] for f in report["layout"][0]["findings"]], ["line_through_text"])
            self.assertFalse(report["pdf"]["older_than_pptx"])
            self.assertTrue((work / "review/slides/slide-02.png").exists())

            older = (work / "deck.pptx").stat().st_mtime - 60
            os.utime(work / "deck.pdf", (older, older))
            stale = subprocess.run(base + ["--pdf", str(work / "deck.pdf"), "--out", str(work / "stale")],
                                   capture_output=True, timeout=120)
            self.assertIn("older than the PPTX", stale.stderr.decode("utf-8", errors="replace"))
            self.assertTrue(json.loads((work / "stale/report.json").read_text(encoding="utf-8"))["pdf"]["older_than_pptx"])

            lines_pdf(work / "three.pdf", pages=3)                                    # not exported from this deck
            wrong = subprocess.run(base + ["--pdf", str(work / "three.pdf"), "--out", str(work / "wrong")],
                                   capture_output=True, timeout=120)
            self.assertEqual(wrong.returncode, 1)
            report = json.loads((work / "wrong/report.json").read_text(encoding="utf-8"))
            self.assertEqual({c["name"]: c["status"] for c in report["checks"]}["render"], "failed")

    def test_zoom_quarters_and_edit_summary_by_kind(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
            for text in ("第一页", "第二页"):
                prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = text
            prs.save(work / "deck.pptx")
            lines_pdf(work / "deck.pdf", pages=2)
            cmd = [sys.executable, str(SCRIPTS / "preview.py"), "pptx", str(work / "deck.pptx"), "--edit-test",
                   "--soffice", str(work / "no-such-soffice.exe"), "--pdf", str(work / "deck.pdf"), "--zoom",
                   "--out", str(work / "review")]
            result = subprocess.run(cmd, capture_output=True, timeout=120)
            self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
            out = result.stdout.decode("utf-8")
            self.assertIn("2 targets (text 2); this deck has no native chart or table", out)
            report = json.loads((work / "review/report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["edit_test"]["targets_by_kind"], {"text": 2})
            quarters = sorted(p.name for p in (work / "review/zoom").glob("*.png"))
            self.assertEqual(len(quarters), 8)
            with Image.open(work / "review/zoom/slide-01-1.png") as quarter:
                self.assertEqual(quarter.size, (1760, 990))                      # 0.55 of 2x 1600

    def test_failed_export_is_reported_not_hidden(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "标题"
            prs.save(work / "deck.pptx")
            real_export, real_argv = preview.export_pdf, sys.argv

            def no_app(pptx, target, app):
                raise RuntimeError(f"{app} is not installed")
            preview.export_pdf = no_app
            sys.argv = ["preview.py", "pptx", str(work / "deck.pptx"), "--export", "wps", "--out", str(work / "review")]
            try:
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    preview.main()
            finally:
                preview.export_pdf, sys.argv = real_export, real_argv
            report = json.loads((work / "review/report.json").read_text(encoding="utf-8"))
            checks = {c["name"]: c["status"] for c in report["checks"]}
            self.assertEqual(checks["render"], "failed")
            self.assertEqual(checks["layout"], "not_run")
            self.assertIn("export with wps failed: wps is not installed", report["render_note"])


qn_ea = "{http://schemas.openxmlformats.org/drawingml/2006/main}ea"
qn_effect = "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst"


def soffice_or_none():
    try:
        return preview.find_soffice(None)
    except SystemExit:
        return None


@unittest.skipUnless(soffice_or_none() and preview.pymupdf_or_none(), "LibreOffice and PyMuPDF are needed to render")
class PptxLayoutTests(unittest.TestCase):
    def test_layout_lint_on_rendered_pptx(self):
        def box(slide, x, y, w, text, size, wrap=True):
            tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1))
            tb.text_frame.word_wrap = wrap
            run = tb.text_frame.paragraphs[0].add_run()
            run.text, run.font.size = text, Pt(size)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        bad = prs.slides.add_slide(prs.slide_layouts[6])
        box(bad, 0.8, 0.6, 8, "重叠的第一段标题文字", 32)
        box(bad, 1.1, 0.75, 8, "重叠的第二段标题文字", 32)
        box(bad, 11.5, 2.5, 4, "这一行文字跑出了页面右边", 24, wrap=False)
        box(bad, 0.8, 4.0, 2.2, "一二三四五六七八九十十", 28)
        box(bad, 6.0, 4.0, 4, "字号太小的说明文字", 7)
        box(bad, 6.0, 5.0, 4, "TODO 补充数据", 18)
        clean = prs.slides.add_slide(prs.slide_layouts[6])
        box(clean, 0.8, 0.6, 10, "干净的页面", 36)
        box(clean, 0.8, 2.0, 10, "正文不重叠，也不跑出页面。", 20)
        table = prs.slides.add_slide(prs.slide_layouts[6])  # one-character number cells are not orphan lines
        rows = [("门店", "服务", "时段", "已约"), ("甲门店", "冷链配送、仓储", "2", "2"), ("乙门店", "冷链配送", "0.5", "0")]
        grid = table.shapes.add_table(len(rows), 4, Inches(0.8), Inches(1), Inches(11), Inches(2)).table
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
        lines = prs.slides.add_slide(prs.slide_layouts[6])
        box(lines, 0.8, 0.6, 5, "被线穿过的标签", 20)
        lines.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(0.87), Inches(6), Inches(0.87))
        lines.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(2.27), Inches(6), Inches(2.27))
        label = lines.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(3), Inches(0.55))
        label.fill.solid()                                                     # painted over the line: hidden
        label.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        label.text_frame.text = "盖住连线的标签"
        label.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        box(lines, 0.8, 3.4, 5, "删除线的旧价格", 20)
        lines.shapes[-1].text_frame.paragraphs[0].runs[0].font._rPr.set("strike", "sngStrike")
        box(lines, 0.8, 4.6, 5, "下划线的说明", 20)
        lines.shapes[-1].text_frame.paragraphs[0].runs[0].font.underline = True
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "layout.pptx"
            prs.save(src)
            layout = []
            preview.pptx_to_pngs(src, Path(tmp) / "out", soffice_or_none(), 800, layout)
            kinds = {f["kind"] for f in layout[0]["findings"]}
            for kind in ("text_overlap", "text_outside_slide", "orphan_line", "font_too_small", "placeholder"):
                self.assertIn(kind, kinds)
            self.assertEqual(layout[1]["findings"], [])
            self.assertEqual(layout[2]["findings"], [])
            self.assertEqual([(f["kind"], f["severity"], f["text"]) for f in layout[3]["findings"]],
                             [("line_through_text", "warning", "被线穿过的标签")])

    def test_orphans_text_past_its_box_and_off_centre_frame(self):
        def box(slide, x, y, w, h, text, size):
            tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tb.text_frame.word_wrap = True
            for i, para in enumerate(text.split("\n")):
                p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
                run = p.add_run()
                run.text, run.font.size = para, Pt(size)

        def card(slide, x, y, w, h):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0xE8, 0xF1, 0xFA)
            shape.line.fill.background()
            shape.shadow.inherit = False                       # LibreOffice still paints the theme shadow ...
            shape._element.remove(shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style"))
            return shape                                       # ... unless the style reference is gone
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        orphans = prs.slides.add_slide(prs.slide_layouts[6])
        box(orphans, 0.8, 1.0, 2.2, 2, "一二三四五六七八九十十；", 28)                # "十；": punctuation aside
        box(orphans, 6.0, 1.0, 2.2, 3, "一二三四五六七八九十十\n下一段文字", 28)       # a lone line before a break
        boxes = prs.slides.add_slide(prs.slide_layouts[6])
        card(boxes, 0.8, 1.0, 4.0, 0.8)
        box(boxes, 1.0, 1.1, 3.6, 0.6, "这段说明比底框长得多，会一直换行，最后从卡片底边跑出去，落在卡片外面。", 18)
        card(boxes, 6.0, 1.0, 5.0, 1.6)
        box(boxes, 6.2, 1.1, 4.6, 1.2, "这段说明完全落在卡片里面。", 18)
        for i, label in enumerate(("第一格", "第二格", "第三格")):                    # stacked cells, one line each
            card(boxes, 0.8, 3.5 + i * 0.4, 3.0, 0.4)
            box(boxes, 0.9, 3.53 + i * 0.4, 2.8, 0.3, label, 16)
        for n in range(3):                                                             # content shifted right
            page = prs.slides.add_slide(prs.slide_layouts[6])
            card(page, 1.1, 2.0, 11.7, 3.0)
            box(page, 1.3, 2.2, 6, 0.6, f"第 {n + 1} 页正文", 20)
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "boxes.pptx"
            prs.save(src)
            layout = []
            preview.pptx_to_pngs(src, Path(tmp) / "out", soffice_or_none(), 800, layout)
        found = [f["kind"] for f in layout[0]["findings"]]
        self.assertEqual(found, ["orphan_line", "orphan_line"])
        self.assertEqual([f["kind"] for f in layout[1]["findings"]], ["text_outside_box"])
        self.assertTrue(layout[1]["findings"][0]["text"].startswith("这段说明比底框长得多"))
        frame = preview.deck_frame([(item["slide"], item["frame"]) for item in layout], "pt", 2, 12)
        self.assertEqual(frame["slides"], [3, 4, 5])
        self.assertAlmostEqual(frame["left"], 79.2, delta=1)
        self.assertAlmostEqual(frame["right"], 38.4, delta=1)


class FrameAndHintTests(unittest.TestCase):
    def test_deck_frame_needs_a_shared_asymmetric_frame(self):
        wide = lambda left, right: {"left": left, "right": right, "span": 960 - left - right, "width": 960}
        shifted = [(n, wide(79, 38)) for n in range(2, 6)] + [(1, {"left": 480, "right": 25, "span": 455, "width": 960})]
        self.assertEqual(preview.deck_frame(shifted, "pt", 2, 12)["slides"], [2, 3, 4, 5])
        self.assertIsNone(preview.deck_frame([(n, wide(45, 45)) for n in range(1, 6)], "pt", 2, 12))  # symmetric
        self.assertIsNone(preview.deck_frame([(1, wide(79, 38)), (2, wide(79, 38))], "pt", 2, 12))  # too few
        self.assertIsNone(preview.deck_frame([(1, None), (2, None), (3, None)], "pt", 2, 12))

    def test_text_heavy_slide_hint_only_in_visual_first(self):
        leaves = [{"text": "社区活动" * 20, "rect": None, "groups": []}, {"text": "报名登记" * 20, "rect": None, "groups": []}]
        kinds = [f["kind"] for f in preview.text_density(leaves, [], True)["findings"]]
        self.assertIn("text_heavy_slide", kinds)
        self.assertEqual(preview.text_density(leaves, [], False)["findings"], [])

    def test_lone_character_ignores_punctuation_but_not_digits(self):
        self.assertTrue(preview.lone_character("间；"))
        self.assertTrue(preview.lone_character("在"))
        self.assertFalse(preview.lone_character("8"))
        self.assertFalse(preview.lone_character("经销商"))


if __name__ == "__main__":
    unittest.main()
