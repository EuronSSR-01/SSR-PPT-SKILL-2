"""Small python-pptx helpers with the defaults this skill relies on. Building blocks, not layouts.

    import sys; sys.path.insert(0, "<skill dir>/scripts")
    from pptx_kit import Deck

    deck = Deck()                               # 1280x720 px canvas on a 13.333x7.5 in slide
    s = deck.slide(bg="FBF9F4")
    deck.text(s, 80, 96, 1000, 60, "生鲜与餐饮占据近六成订单", size=44, font="Noto Serif SC", bold=True)
    deck.bar_chart(s, 80, 220, 790, 380, ["生鲜", "餐饮", "百货"], [34, 24, 15], highlight={0, 1})
    deck.notes(s, "讲稿与来源")
    deck.save("deck.pptx", theme_fonts=("Noto Serif SC", "Noto Sans SC"))

All positions and sizes are canvas px (default 1280x720) and are converted with the real slide size,
so a 1280 px HTML layout maps 1:1. Font sizes (size=) and line widths are canvas px too, not pt: on the
default 13.333 in wide slide 1 px = 0.75 pt, so 12 pt body text is size=16. For other slide sizes,
pt = px * slide width (in) * 72 / canvas width; check the real pt in the saved file.
Defaults that prevent known problems:
- every run gets both Latin and East Asian fonts (python-pptx alone writes only Latin);
- text boxes have zero insets and never auto-shrink;
- shapes have no shadow in any renderer (no theme style reference) and no default outline;
- bar/column charts keep one real data series, highlight points by colouring them, put labels
  outside the bars and drop legend, gridlines and the value axis;
- tables use no built-in style: hairline rows, quiet header, numbers right-aligned on request;
- pictures are cropped to fill their box instead of being stretched;
- icons from assets/icons become native line shapes (recolour in PowerPoint/WPS, sharp at any size);
- SVG illustrations are rasterised at twice the box size with PyMuPDF and inserted as pictures.
With a company template, pass Deck(template=...) and pick its layouts; theme fonts are left alone.
"""
import importlib.util
import io
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.image import Image
from pptx.util import Emu, Pt

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # built-in "No Style, No Grid"
ICON_DIR = Path(__file__).resolve().parent.parent / "assets" / "icons"
SVG_NS = "{http://www.w3.org/2000/svg}"


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


# ---------------------------------------------------------------- SVG geometry -> DrawingML
# Enough SVG for line icons: path (all commands, arcs become cubic curves), circle, ellipse, rect,
# line, polyline and polygon. Transforms, styles and text are not supported.

_TOKEN = re.compile(r"[MmLlHhVvCcSsQqTtAaZz]|[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?")


def _arc_to_cubics(x1, y1, rx, ry, phi, large, sweep, x2, y2):
    """SVG endpoint arc -> list of cubic segments (c1x, c1y, c2x, c2y, x, y)."""
    if rx == 0 or ry == 0 or (x1 == x2 and y1 == y2):
        return [(x1, y1, x2, y2, x2, y2)]
    rx, ry = abs(rx), abs(ry)
    cos_p, sin_p = math.cos(math.radians(phi)), math.sin(math.radians(phi))
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    xp, yp = cos_p * dx + sin_p * dy, -sin_p * dx + cos_p * dy
    scale = (xp * xp) / (rx * rx) + (yp * yp) / (ry * ry)
    if scale > 1:
        rx, ry = rx * math.sqrt(scale), ry * math.sqrt(scale)
    num = rx * rx * ry * ry - rx * rx * yp * yp - ry * ry * xp * xp
    den = rx * rx * yp * yp + ry * ry * xp * xp
    coef = math.sqrt(max(0.0, num / den)) * (-1 if large == sweep else 1)
    cxp, cyp = coef * rx * yp / ry, -coef * ry * xp / rx
    cx = cos_p * cxp - sin_p * cyp + (x1 + x2) / 2
    cy = sin_p * cxp + cos_p * cyp + (y1 + y2) / 2

    def angle(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a

    t1 = angle(1, 0, (xp - cxp) / rx, (yp - cyp) / ry)
    dt = angle((xp - cxp) / rx, (yp - cyp) / ry, (-xp - cxp) / rx, (-yp - cyp) / ry)
    if not sweep and dt > 0:
        dt -= 2 * math.pi
    elif sweep and dt < 0:
        dt += 2 * math.pi
    parts = max(1, math.ceil(abs(dt) / (math.pi / 2) - 1e-9))
    step = dt / parts
    k = 4 / 3 * math.tan(step / 4)
    out = []

    def point(t):
        x, y = rx * math.cos(t), ry * math.sin(t)
        return cos_p * x - sin_p * y + cx, sin_p * x + cos_p * y + cy

    def deriv(t):
        x, y = -rx * math.sin(t), ry * math.cos(t)
        return cos_p * x - sin_p * y, sin_p * x + cos_p * y

    t = t1
    for _ in range(parts):
        p0, p3 = point(t), point(t + step)
        d0, d3 = deriv(t), deriv(t + step)
        out.append((p0[0] + k * d0[0], p0[1] + k * d0[1], p3[0] - k * d3[0], p3[1] - k * d3[1], p3[0], p3[1]))
        t += step
    out[-1] = out[-1][:4] + (x2, y2)
    return out


def svg_path_to_subpaths(d):
    """Parse an SVG path into absolute subpaths of ("M", x, y), ("L", x, y), ("C", 6 numbers), ("Z",)."""
    tokens = _TOKEN.findall(d)
    subpaths, current = [], []
    i, cmd = 0, None
    x = y = sx = sy = 0.0
    last_ctrl, last_cmd = None, None

    def nums(n):
        nonlocal i
        values = [float(t) for t in tokens[i:i + n]]
        if len(values) < n or any(t.isalpha() for t in tokens[i:i + n]):
            raise ValueError(f"bad path data near {' '.join(tokens[i:i + n])!r}")
        i += n
        return values

    while i < len(tokens):
        if tokens[i].isalpha():
            cmd = tokens[i]
            i += 1
        elif cmd is None:
            raise ValueError("path data must start with a command")
        rel = cmd.islower()
        c = cmd.upper()
        ox, oy = (x, y) if rel else (0.0, 0.0)
        if c == "Z":
            current.append(("Z",))
            x, y = sx, sy
            subpaths.append(current)
            current = []
            last_ctrl, last_cmd = None, "Z"
            continue
        if c == "M":
            if current:
                subpaths.append(current)
            px, py = nums(2)
            x, y = ox + px, oy + py
            sx, sy = x, y
            current = [("M", x, y)]
            cmd = "l" if rel else "L"  # extra pairs after a moveto are lineto
            last_ctrl, last_cmd = None, "M"
            continue
        if not current:
            current = [("M", x, y)]
            sx, sy = x, y
        if c == "L":
            px, py = nums(2)
            x, y = ox + px, oy + py
            current.append(("L", x, y))
            last_ctrl = None
        elif c == "H":
            (px,) = nums(1)
            x = ox + px
            current.append(("L", x, y))
            last_ctrl = None
        elif c == "V":
            (py,) = nums(1)
            y = oy + py
            current.append(("L", x, y))
            last_ctrl = None
        elif c in "CS":
            if c == "C":
                x1, y1, x2, y2, px, py = nums(6)
                x1, y1 = ox + x1, oy + y1
            else:
                x2, y2, px, py = nums(4)
                if last_ctrl and last_cmd in "CS":
                    x1, y1 = 2 * x - last_ctrl[0], 2 * y - last_ctrl[1]
                else:
                    x1, y1 = x, y
            x2, y2, px, py = ox + x2, oy + y2, ox + px, oy + py
            current.append(("C", x1, y1, x2, y2, px, py))
            last_ctrl = (x2, y2)
            x, y = px, py
        elif c in "QT":
            if c == "Q":
                qx, qy, px, py = nums(4)
                qx, qy = ox + qx, oy + qy
            else:
                px, py = nums(2)
                if last_ctrl and last_cmd in "QT":
                    qx, qy = 2 * x - last_ctrl[0], 2 * y - last_ctrl[1]
                else:
                    qx, qy = x, y
            px, py = ox + px, oy + py
            current.append(("C", x + 2 / 3 * (qx - x), y + 2 / 3 * (qy - y),
                            px + 2 / 3 * (qx - px), py + 2 / 3 * (qy - py), px, py))
            last_ctrl = (qx, qy)
            x, y = px, py
        elif c == "A":
            rx, ry, phi = nums(3)
            flags = []
            while len(flags) < 2:  # flags may be packed without separators, e.g. "011"
                token = tokens[i]
                flags.append(int(token[0]))
                rest = token[1:]
                if rest:
                    tokens[i] = rest
                else:
                    i += 1
            px, py = nums(2)
            px, py = ox + px, oy + py
            for seg in _arc_to_cubics(x, y, rx, ry, phi, flags[0], flags[1], px, py):
                current.append(("C",) + seg)
            x, y = px, py
            last_ctrl = None
        else:
            raise ValueError(f"unsupported path command {cmd!r}")
        last_cmd = c
    if current:
        subpaths.append(current)
    return subpaths


def _ellipse(cx, cy, rx, ry):
    k = 0.5522847498
    return [("M", cx + rx, cy),
            ("C", cx + rx, cy + k * ry, cx + k * rx, cy + ry, cx, cy + ry),
            ("C", cx - k * rx, cy + ry, cx - rx, cy + k * ry, cx - rx, cy),
            ("C", cx - rx, cy - k * ry, cx - k * rx, cy - ry, cx, cy - ry),
            ("C", cx + k * rx, cy - ry, cx + rx, cy - k * ry, cx + rx, cy), ("Z",)]


def svg_element_subpaths(el):
    """Subpaths for one SVG shape element; empty list for anything else."""
    tag = el.tag.replace(SVG_NS, "")
    f = lambda name, default=0.0: float(el.get(name, default))
    if tag == "path":
        return svg_path_to_subpaths(el.get("d", ""))
    if tag == "circle":
        return [_ellipse(f("cx"), f("cy"), f("r"), f("r"))]
    if tag == "ellipse":
        return [_ellipse(f("cx"), f("cy"), f("rx"), f("ry"))]
    if tag == "line":
        return [[("M", f("x1"), f("y1")), ("L", f("x2"), f("y2"))]]
    if tag in ("polyline", "polygon"):
        values = [float(v) for v in re.findall(r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?", el.get("points", ""))]
        pts = list(zip(values[0::2], values[1::2]))
        if not pts:
            return []
        sub = [("M",) + pts[0]] + [("L",) + p for p in pts[1:]]
        return [sub + [("Z",)] if tag == "polygon" else sub]
    if tag == "rect":
        x, y, w, h = f("x"), f("y"), f("width"), f("height")
        rx = el.get("rx", el.get("ry"))
        ry = el.get("ry", el.get("rx"))
        rx = min(float(rx), w / 2) if rx is not None else 0.0
        ry = min(float(ry), h / 2) if ry is not None else 0.0
        if not rx or not ry:
            return [[("M", x, y), ("L", x + w, y), ("L", x + w, y + h), ("L", x, y + h), ("Z",)]]
        d = (f"M{x + rx},{y} H{x + w - rx} A{rx},{ry} 0 0 1 {x + w},{y + ry} V{y + h - ry} "
             f"A{rx},{ry} 0 0 1 {x + w - rx},{y + h} H{x + rx} A{rx},{ry} 0 0 1 {x},{y + h - ry} "
             f"V{y + ry} A{rx},{ry} 0 0 1 {x + rx},{y} Z")
        return svg_path_to_subpaths(d)
    return []


def icon_path(name):
    """Path of a bundled icon (assets/icons/<name>.svg); see assets/icons/INDEX.md for names."""
    path = ICON_DIR / f"{name}.svg"
    if not path.exists():
        raise KeyError(f"icon {name!r} not found in {ICON_DIR}; see INDEX.md there")
    return path


def svg_geometry(svg):
    """(viewBox width, viewBox height, [(subpaths, filled)]) for an SVG file path or SVG text."""
    text = Path(svg).read_text(encoding="utf-8") if not str(svg).lstrip().startswith("<") else str(svg)
    root = ElementTree.fromstring(text[text.index("<svg"):])
    box = [float(v) for v in re.split(r"[\s,]+", root.get("viewBox", "0 0 24 24").strip())]
    root_fill = root.get("fill", "black")
    shapes = []
    for el in root.iter():
        subpaths = svg_element_subpaths(el)
        if subpaths:
            fill = el.get("fill", root_fill)
            shapes.append(([[(c[0],) + tuple(v - (box[0] if j % 2 == 0 else box[1]) for j, v in enumerate(c[1:]))
                             for c in sub] for sub in subpaths], fill not in ("none", "transparent")))
    return box[2], box[3], shapes


def plain(shape):
    """No shadow in any renderer: an empty effect list plus no theme style reference (LibreOffice still
    paints the style's shadow when only the effect list is empty). Fill and line are set explicitly."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def set_ea(element, typeface):
    """Give every Latin font under element a matching East Asian font (a:ea right after a:latin)."""
    for latin in element.iter(qn("a:latin")):
        ea = latin.getnext()
        if ea is None or ea.tag != qn("a:ea"):
            ea = OxmlElement("a:ea")
            latin.addnext(ea)
        ea.set("typeface", typeface)


class Deck:
    def __init__(self, template=None, width=1280, height=720):
        self.template = template
        self.prs = Presentation(template) if template else Presentation()
        if not template:
            self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in
        self.k = self.prs.slide_width / width  # EMU per canvas px
        ratio = (self.prs.slide_width / self.prs.slide_height) / (width / height)
        if abs(ratio - 1) > 0.01:
            raise ValueError(f"canvas {width}x{height} does not match the slide aspect ratio; pass matching width/height")

    def px(self, value):
        return Emu(round(value * self.k))

    def pt(self, size_px):
        return Pt(size_px * self.k / 12700)

    # ------------------------------------------------------------ slides
    def slide(self, layout=None, bg=None):
        """New slide. layout: a layout name from the template; default is the layout with fewest placeholders."""
        layouts = list(self.prs.slide_layouts)
        if layout is None:
            chosen = min(layouts, key=lambda l: len(l.placeholders))
        else:
            chosen = next((l for l in layouts if l.name == layout), None)
            if chosen is None:
                raise KeyError(f"layout {layout!r} not found; available: {[l.name for l in layouts]}")
        slide = self.prs.slides.add_slide(chosen)
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(bg)
        return slide

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text

    # ------------------------------------------------------------ text
    def _style_run(self, run, font=None, ea=None, size=None, color=None, bold=None, italic=None):
        if font:
            run.font.name = font
            set_ea(run._r, ea or font)
        if size:
            run.font.size = self.pt(size)
        if color:
            run.font.color.rgb = rgb(color)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic

    def text(self, slide, x, y, w, h, content, size=22, font="Microsoft YaHei", ea=None, color="1A1A1A",
             bold=False, align="left", valign="top", line_spacing=None, wrap=True, space_after=0):
        """Text box. content: a string ("\\n" starts a paragraph) or a list of paragraphs, each a string or
        a list of runs: ("text", {"color": ..., "bold": ..., "size": ..., "font": ...})."""
        box = slide.shapes.add_textbox(self.px(x), self.px(y), self.px(w), self.px(h))
        tf = box.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[valign]
        paragraphs = content.split("\n") if isinstance(content, str) else content
        for index, para in enumerate(paragraphs):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            if line_spacing:
                p.line_spacing = line_spacing
            p.space_after = self.pt(space_after) if space_after else 0
            runs = [(para, {})] if isinstance(para, str) else para
            for text, opts in runs:
                run = p.add_run()
                run.text = text
                self._style_run(run, font=opts.get("font", font), ea=opts.get("ea", ea), size=opts.get("size", size),
                                color=opts.get("color", color), bold=opts.get("bold", bold), italic=opts.get("italic"))
        return box

    # ------------------------------------------------------------ shapes
    def rect(self, slide, x, y, w, h, fill=None, line=None, line_width=1, radius=None):
        """Rectangle without inherited shadow; radius (px) gives rounded corners."""
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, self.px(x), self.px(y), self.px(w), self.px(h))
        if radius:
            shape.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = self.pt(line_width)
        else:
            shape.line.fill.background()
        return plain(shape)

    def ellipse(self, slide, cx, cy, r, fill=None, line=None, line_width=1):
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, self.px(cx - r), self.px(cy - r), self.px(2 * r), self.px(2 * r))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = self.pt(line_width)
        else:
            shape.line.fill.background()
        return plain(shape)

    def line(self, slide, x1, y1, x2, y2, color="1A1A1A", width=1, arrow=False, dash=False):
        """Straight connector; arrow=True adds an arrowhead at (x2, y2)."""
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.px(x1), self.px(y1), self.px(x2), self.px(y2))
        conn.line.color.rgb = rgb(color)
        conn.line.width = self.pt(width)
        if dash:
            conn.line.dash_style = MSO_LINE.DASH
        if arrow:
            tail = OxmlElement("a:tailEnd")
            tail.set("type", "triangle")
            conn.line._get_or_add_ln().append(tail)
        return conn

    def picture(self, slide, path, x, y, w, h, fit="cover"):
        """Picture cropped to fill the box (fit="cover") or shrunk to fit inside it (fit="contain")."""
        img_w, img_h = Image.from_file(str(path)).size
        box_ratio, img_ratio = w / h, img_w / img_h
        if fit == "contain":
            if img_ratio > box_ratio:
                h2 = w / img_ratio
                return slide.shapes.add_picture(str(path), self.px(x), self.px(y + (h - h2) / 2), self.px(w), self.px(h2))
            w2 = h * img_ratio
            return slide.shapes.add_picture(str(path), self.px(x + (w - w2) / 2), self.px(y), self.px(w2), self.px(h))
        pic = slide.shapes.add_picture(str(path), self.px(x), self.px(y), self.px(w), self.px(h))
        if img_ratio > box_ratio:
            pic.crop_left = pic.crop_right = (1 - box_ratio / img_ratio) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - img_ratio / box_ratio) / 2
        return pic

    # ------------------------------------------------------------ icons and SVG illustrations
    def icon(self, slide, name, x, y, size=32, color="1A1A1A", stroke=2.0, fill=None):
        """Bundled line icon (assets/icons/<name>.svg, or any SVG file path) as one native shape.
        size: px box; stroke: line width in icon units (Lucide draws 2 on a 24 grid); fill: colour for
        elements the SVG fills. The shape can be recoloured, resized and moved like any other shape."""
        source = Path(name) if str(name).lower().endswith(".svg") else icon_path(name)
        vb_w, vb_h, shapes = svg_geometry(source)
        scale = 1000  # path coordinates are integers in a vb_w*1000 x vb_h*1000 space
        paths = []
        for subpaths, filled in shapes:
            path = OxmlElement("a:path")
            path.set("w", str(round(vb_w * scale)))
            path.set("h", str(round(vb_h * scale)))
            if not filled:
                path.set("fill", "none")
            for sub in subpaths:
                for c in sub:
                    if c[0] == "Z":
                        path.append(OxmlElement("a:close"))
                        continue
                    tag = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo"}[c[0]]
                    node = OxmlElement(tag)
                    for px_, py_ in zip(c[1::2], c[2::2]):
                        pt = OxmlElement("a:pt")
                        pt.set("x", str(round(px_ * scale)))
                        pt.set("y", str(round(py_ * scale)))
                        node.append(pt)
                    path.append(node)
            paths.append(path)
        h = size * vb_h / vb_w
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, self.px(x), self.px(y), self.px(size), self.px(h))
        shape.name = f"Icon {Path(source).stem}"
        sp_pr = shape._element.spPr
        geom = OxmlElement("a:custGeom")
        for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
            geom.append(OxmlElement(tag))
        rect = OxmlElement("a:rect")
        for k, v in (("l", "l"), ("t", "t"), ("r", "r"), ("b", "b")):
            rect.set(k, v)
        geom.append(rect)
        path_lst = OxmlElement("a:pathLst")
        for path in paths:
            path_lst.append(path)
        geom.append(path_lst)
        sp_pr.replace(sp_pr.find(qn("a:prstGeom")), geom)
        if any(filled for _, filled in shapes):
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill or color)
        else:
            shape.fill.background()
        shape.line.color.rgb = rgb(color)
        shape.line.width = self.pt(stroke * size / vb_w)
        ln = shape.line._get_or_add_ln()
        ln.set("cap", "rnd")
        if ln.find(qn("a:round")) is None:
            ln.append(OxmlElement("a:round"))
        return plain(shape)

    def svg_picture(self, slide, svg, x, y, w, h, out=None, fit="contain", density=2):
        """SVG illustration (file path or SVG text) rasterised with PyMuPDF at density x the box size and
        inserted as a picture. Keep the SVG source in the build folder; edit it and rerun to change
        the drawing. out: PNG path to write (default: next to the SVG, or svg-<n>.png)."""
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz  # older PyMuPDF
        text = str(svg)
        is_text = text.lstrip().startswith("<")
        data = text.encode("utf-8") if is_text else Path(svg).read_bytes()
        doc = fitz.open(stream=data, filetype="svg")
        page = doc[0]
        zoom = max(self.px(w) / 12700 / page.rect.width, self.px(h) / 12700 / page.rect.height) * 96 / 72 * density
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
        if out is None:
            out = Path(svg).with_suffix(".png") if not is_text else Path(f"svg-{len(slide.shapes) + 1}.png")
        pix.save(str(out))
        doc.close()
        return self.picture(slide, out, x, y, w, h, fit=fit)

    # ------------------------------------------------------------ charts and tables
    def bar_chart(self, slide, x, y, w, h, categories, values, highlight=(), horizontal=True,
                  color="E8590C", base="CBC5B8", series_name="值", number_format="0",
                  font="Microsoft YaHei", size=16, label_color="666158", axis_color="1A1A1A",
                  gap_width=60, max_value=None, min_value=None):
        """Native bar (horizontal) or column chart with one real series; highlight = indices to colour.
        The value range always includes zero and leaves room for the outside labels on each side that has
        data; pass max_value / min_value to fix either end."""
        numeric = [v for v in values if v is not None]
        if not numeric:
            raise ValueError("bar_chart needs at least one numeric value")
        data = CategoryChartData()  # values are passed through unchanged; only the axis range is computed
        data.categories = list(categories)
        data.add_series(series_name, list(values))
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        frame = slide.shapes.add_chart(kind, self.px(x), self.px(y), self.px(w), self.px(h), data)
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False
        chart.font.name = font
        chart.font.size = self.pt(size)
        plot = chart.plots[0]
        plot.gap_width = gap_width
        plot.vary_by_categories = False
        series = plot.series[0]
        series.invert_if_negative = False  # otherwise some renderers paint negative bars white (invisible)
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(base)
        for i in highlight:
            point = series.points[i]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(color)
        for dpt in series._element.findall(qn("c:dPt")):  # each point needs its own flag too, right after c:idx
            if dpt.find(qn("c:invertIfNegative")) is None:
                flag = OxmlElement("c:invertIfNegative")
                flag.set("val", "0")
                dpt.find(qn("c:idx")).addnext(flag)
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = number_format
        labels.number_format_is_linked = False
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
        labels.font.size = self.pt(size)
        labels.font.color.rgb = rgb(label_color)
        value_axis, category_axis = chart.value_axis, chart.category_axis
        value_axis.visible = False
        value_axis.has_major_gridlines = False
        value_axis.has_minor_gridlines = False
        low, high = min(0, min(numeric)), max(0, max(numeric))
        pad = ((high - low) or 1) * 0.18  # room for labels beyond the bar ends
        top = max_value if max_value is not None else (high + pad if high > 0 else 0 if low < 0 else 1)
        bottom = min_value if min_value is not None else (low - pad if low < 0 else 0)
        if not bottom < top:
            raise ValueError(f"value range [{bottom}, {top}] is empty; check min_value / max_value")
        value_axis.maximum_scale, value_axis.minimum_scale = top, bottom
        category_axis.has_major_gridlines = False
        category_axis.format.line.fill.background()
        category_axis.tick_labels.font.size = self.pt(size)
        category_axis.tick_labels.font.color.rgb = rgb(axis_color)
        category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # beside the plot, not on the zero line
        if horizontal:
            category_axis.reverse_order = True  # first category on top
        set_ea(chart._chartSpace, font)
        return frame

    def table(self, slide, x, y, w, rows, col_widths=None, row_height=48, size=18, font="Microsoft YaHei",
              color="1A1A1A", header=True, header_color="5B6878", line_color="D9DEE5", header_line="1A1A1A",
              align=None, emphasis=None):
        """Native table without a built-in style. rows: list of lists of strings (first row = header).
        align: per-column "left" / "right" / "center". emphasis: {(row, col): "HEX"} colours those cells' text."""
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        widths = col_widths or [w / n_cols] * n_cols
        frame = slide.shapes.add_table(n_rows, n_cols, self.px(x), self.px(y), self.px(w), self.px(row_height * n_rows))
        table = frame.table
        tbl_pr = table._tbl.tblPr
        tbl_pr.set("firstRow", "0")
        tbl_pr.set("bandRow", "0")
        style = tbl_pr.find(qn("a:tableStyleId"))
        if style is None:
            style = OxmlElement("a:tableStyleId")
            tbl_pr.append(style)
        style.text = NO_TABLE_STYLE
        for c, cw in enumerate(widths):
            table.columns[c].width = self.px(cw)
        emphasis = emphasis or {}
        for r in range(n_rows):
            table.rows[r].height = self.px(row_height)
            is_header = header and r == 0
            for c in range(n_cols):
                cell = table.cell(r, c)
                cell.fill.background()
                cell.margin_left = cell.margin_right = self.px(6)
                cell.margin_top = cell.margin_bottom = 0
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                text = rows[r][c] if c < len(rows[r]) else ""
                p = cell.text_frame.paragraphs[0]
                p.alignment = ALIGN[(align or ["left"] * n_cols)[c]]
                run = p.add_run()
                run.text = text
                self._style_run(run, font=font, size=size * 0.8 if is_header else size,
                                color=header_color if is_header else emphasis.get((r, c), color),
                                bold=(r, c) in emphasis)
                self._bottom_border(cell, header_line if is_header else line_color, 1.5 if is_header else 0.75)
        return frame

    @staticmethod
    def _bottom_border(cell, color, width_pt):
        tc_pr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            old = tc_pr.find(qn(tag))
            if old is not None:
                tc_pr.remove(old)
        borders = []
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = OxmlElement(tag)
            if tag == "a:lnB":
                ln.set("w", str(int(width_pt * 12700)))
                fill = OxmlElement("a:solidFill")
                clr = OxmlElement("a:srgbClr")
                clr.set("val", color.lstrip("#").upper())
                fill.append(clr)
                ln.append(fill)
            else:
                ln.append(OxmlElement("a:noFill"))
            borders.append(ln)
        for ln in reversed(borders):  # borders must precede the cell fill
            tc_pr.insert(0, ln)

    # ------------------------------------------------------------ save
    def save(self, path, theme_fonts=None):
        """Save; theme_fonts=(heading, body) also sets the theme's Latin and East Asian fonts.
        Ignored for company templates, whose theme is kept unless the task explicitly asks otherwise."""
        path = Path(path)
        self.prs.save(path)
        if theme_fonts and not self.template:
            spec = importlib.util.spec_from_file_location("set_theme_fonts", Path(__file__).with_name("set_theme_fonts.py"))
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            set_block = module.set_block
            heading, body = theme_fonts
            buf = io.BytesIO()
            with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("ppt/theme/theme") and item.filename.endswith(".xml"):
                        xml = data.decode("utf-8")
                        xml, _ = set_block(xml, "majorFont", heading, heading)
                        xml, _ = set_block(xml, "minorFont", body, body)
                        data = xml.encode("utf-8")
                    zout.writestr(item, data)
            path.write_bytes(buf.getvalue())
        return path
