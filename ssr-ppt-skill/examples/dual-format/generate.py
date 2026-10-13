#!/usr/bin/env python3
"""Small, explicit shared-source example; not a general HTML-to-PPTX converter."""
import argparse
import base64
from html import escape
import json
from pathlib import Path
import subprocess
import sys
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SKILL = HERE.parents[1]
W, H = 1280, 720
BG, INK, MUTED, ACCENT, LIGHT, LINE = "F6F3EC", "18342F", "576B65", "176A58", "E3EAE2", "C8D2C8"
FONT = "Microsoft YaHei"


def inch(px):
    return Inches(px / 96)


def color(hex_value):
    return RGBColor.from_string(hex_value)


def font(run, size, shade=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size * 0.75)
    run.font.color.rgb = color(shade)
    run.font.bold = bold
    ea = OxmlElement("a:ea")
    ea.set("typeface", FONT)
    run._r.get_or_add_rPr().append(ea)


def text(shapes, svg, value, x, y, width, size=24, shade=INK, bold=False):
    lines = str(value).split("\n")
    box = shapes.add_textbox(inch(x), inch(y), inch(width), inch(len(lines)*size*1.4+8))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_top = frame.margin_right = frame.margin_bottom = 0
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_before = p.space_after = 0
        p.line_spacing = 1.4
        run = p.add_run()
        run.text = line
        font(run, size, shade, bold)
        svg.append(f'<text x="{x}" y="{y+size+i*size*1.4}" font-size="{size}" fill="#{shade}" font-weight="{700 if bold else 400}">{escape(line)}</text>')
    return box


def rect(shapes, svg, x, y, w, h, fill, stroke=None):
    shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    for effect in shape._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#{fill}" stroke="#{stroke or fill}"/>')
    return shape


def arrow(shapes, svg, x1, y1, x2, y2, start=None, end=None, arrow_end=True):
    line = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    line.line.color.rgb = color(ACCENT)
    line.line.width = Pt(1.5)
    if arrow_end:
        head = OxmlElement("a:tailEnd")
        head.set("type", "triangle")
        line._element.spPr.get_or_add_ln().append(head)
    for effect in line._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    if start:
        line.begin_connect(start, 3)
    if end:
        line.end_connect(end, 1)
    marker = ' marker-end="url(#arrow)"' if arrow_end else ''
    svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#{ACCENT}" stroke-width="2"{marker}/>')
    return line


def illustration(path):
    # Synthetic crop-test artwork, intentionally neither a real photo nor evidence.
    im = Image.new("RGB", (900, 900), "#E3EAE2")
    draw = ImageDraw.Draw(im)
    for i in range(5):
        pad = 55+i*65
        draw.rounded_rectangle((pad, pad, 900-pad, 900-pad), radius=65, outline="#176A58", width=12)
    draw.ellipse((325, 325, 575, 575), fill="#D6AC62")
    im.save(path)


def generate(out):
    out.mkdir(parents=True, exist_ok=True)
    data = json.loads((HERE / "data.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = inch(W), inch(H)
    prs.core_properties.title = data["title"] + " · 双格式教学样例"
    prs.core_properties.author = "SSR-PPT-SKILL"
    pages, targets = [], []
    def target(shape, kind, group=None, **extra):
        path = [group.shape_id, shape.shape_id] if group else [shape.shape_id]
        targets.append(dict(slide=len(prs.slides), path=path, kind=kind, **extra))
    def page(title, number):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(BG)
        svg = [f'<rect width="1280" height="720" fill="#{BG}"/>']
        text(slide.shapes, svg, "SSR / 教学样例", 80, 40, 550, 16, MUTED)
        text(slide.shapes, svg, f"{number:02d} / 04", 1090, 40, 140, 16, MUTED)
        title_box = text(slide.shapes, svg, title, 80, 105, 1140, 44, bold=True)
        target(title_box, "text")
        text(slide.shapes, svg, "虚构数据  /  截至 " + data["date"], 80, 665, 700, 14, MUTED)
        slide.notes_slide.notes_text_frame.text = data["notice"]
        pages.append(svg)
        return slide, svg

    # 1. Claim, status and an independent image; no decorative statistical chart.
    slide, svg = page("上线进度清楚，\n资源与收益仍需分别确认", 1)
    text(slide.shapes, svg, f'{data["live"]}/{data["planned"]}', 80, 270, 450, 130, ACCENT, True)
    text(slide.shapes, svg, "个计划站点已上线", 85, 470, 540, 28)
    text(slide.shapes, svg, f'申请中：{data["requested_gateways"]} 台网关\n预测节约：{data["forecast_hours"]} 小时／月，待实测', 85, 535, 600, 24, MUTED)
    image_path = out / "illustration.png"
    illustration(image_path)
    pic = slide.shapes.add_picture(str(image_path), inch(790), inch(260), width=inch(350), height=inch(350))
    target(pic, "picture")
    uri = "data:image/png;base64,"+base64.b64encode(image_path.read_bytes()).decode("ascii")
    svg.append(f'<image x="790" y="260" width="350" height="350" href="{uri}"/>')
    text(slide.shapes, svg, "抽象插画 · 非现场照片", 800, 620, 400, 16, MUTED)

    # 2. Semantic flow. Grouped editable labels and real bound connectors.
    slide, svg = page("通过才备货，周报不参与审批", 2)
    text(slide.shapes, svg, "把动作、条件和信息流分开，避免画成没有依据的层级。", 80, 190, 1120, 24, MUTED)
    group = slide.shapes.add_group_shape()
    nodes = []
    for i, item in enumerate(data["process"]):
        x = 80+i*390
        nodes.append(rect(group.shapes, svg, x, 295, 280, 130, LIGHT))
        label = text(group.shapes, svg, item["title"], x+24, 313, 230, 30, ACCENT, True)
        text(group.shapes, svg, item["detail"], x+24, 365, 240, 22)
        if i == 1:
            target(label, "text", group)
    arrow(group.shapes, svg, 360, 360, 470, 360, nodes[0], nodes[1])
    arrow(group.shapes, svg, 750, 360, 860, 360, nodes[1], nodes[2])
    text(group.shapes, svg, "提交", 380, 315, 90, 18, MUTED)
    text(group.shapes, svg, "通过", 772, 315, 90, 18, MUTED)
    # Explicit return route; the path's meaning is carried by its label.
    arrow(group.shapes, svg, 610, 435, 610, 500, arrow_end=False)
    arrow(group.shapes, svg, 610, 500, 220, 500, arrow_end=False)
    arrow(group.shapes, svg, 220, 500, 220, 435)
    text(group.shapes, svg, "不通过 → 门店补充信息", 280, 525, 450, 22, MUTED)
    text(group.shapes, svg, "运营接收周报\n不参与本次审批", 870, 490, 320, 22, MUTED)
    target(group, "move")

    # 3. A comparison matrix made of grouped shapes, not a mandatory native table.
    slide, svg = page("先比较交付约束，再决定预算", 3)
    text(slide.shapes, svg, "不做没有依据的综合排名；三个方案分别保留条件。", 80, 190, 1100, 24, MUTED)
    group = slide.shapes.add_group_shape()
    rows = ["基础报价", "预计交付", "扩展条件"]
    for i, item in enumerate(data["options"]):
        x = 300+i*290
        rect(group.shapes, svg, x, 270, 265, 310, LIGHT, LINE)
        text(group.shapes, svg, item["name"], x+24, 288, 220, 27, ACCENT, True)
        for j, key in enumerate(("cost", "weeks", "condition")):
            value = text(group.shapes, svg, item[key], x+24, 372+j*69, 240, 22 if j == 2 else 28, INK, j == 0)
            if i == 1 and key == "cost":
                target(value, "text", group)
    for j, value in enumerate(rows):
        text(group.shapes, svg, value, 80, 377+j*69, 200, 24, MUTED)
    target(group, "move")
    text(slide.shapes, svg, "所有报价均为虚构。扩展费用不包含在基础报价内。", 80, 610, 1120, 22, MUTED)

    # 4. One standard native chart, a real native table, and visible data conditions.
    slide, svg = page("上线数在增加，收益仍需实测", 4)
    text(slide.shapes, svg, "累计上线站点／个，同一计划范围", 80, 200, 650, 24, MUTED)
    chart_data = CategoryChartData()
    chart_data.categories = data["trend"]["categories"]
    chart_data.add_series("累计上线", data["trend"]["values"])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, inch(65), inch(270), inch(640), inch(310), chart_data)
    chart = chart_shape.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name, chart.font.size = FONT, Pt(16)
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = 0, data["planned"]
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.visible = False
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(16)
    plot = chart.plots[0]
    plot.gap_width = 95
    plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(18)
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = color(ACCENT)
    plot.series[0].format.line.fill.background()
    target(chart_shape, "chart")
    for i, value in enumerate(data["trend"]["values"]):
        x, h = 125+i*185, value*34
        svg.append(f'<rect x="{x}" y="{550-h}" width="94" height="{h}" fill="#{ACCENT}"/>')
        svg.append(f'<text x="{x+47}" y="{532-h}" text-anchor="middle" font-size="26" fill="#{INK}">{value}</text>')
        svg.append(f'<text x="{x+47}" y="587" text-anchor="middle" font-size="22" fill="#{MUTED}">{escape(data["trend"]["categories"][i])}</text>')
    text(slide.shapes, svg, f'{data["forecast_hours"]} 小时／月', 790, 270, 430, 44, ACCENT, True)
    text(slide.shapes, svg, "预测收益 · 尚未验证", 790, 344, 430, 24, MUTED)
    table_shape = slide.shapes.add_table(3, 2, inch(790), inch(430), inch(410), inch(160))
    table = table_shape.table
    values = [["待核事项", "当前状态"], ["资源审批", "申请中"], ["收益测量", "未完成"]]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(LIGHT if r == 0 else BG)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    font(run, 21, INK, r == 0)
            svg.append(f'<rect x="{790+c*205}" y="{430+r*53}" width="205" height="53" fill="#{LIGHT if r == 0 else BG}" stroke="#{LINE}"/>')
            svg.append(f'<text x="{805+c*205}" y="{463+r*53}" font-size="21" fill="#{INK}">{value}</text>')
    target(table_shape, "table", cell=[2, 1])
    text(slide.shapes, svg, "上线不等于收益兑现；预测不并入实际结果。", 80, 615, 1150, 22, MUTED)

    prs.save(out / "deck.pptx")
    # This is a fresh synthetic deck, including its generated notes theme.
    subprocess.run([sys.executable, str(SKILL / "scripts" / "set_theme_fonts.py"), str(out / "deck.pptx"), "--heading", FONT, "--body", FONT, "--all-themes"], check=True)
    css = (SKILL / "assets" / "player.css").read_text(encoding="utf-8")
    js = (SKILL / "assets" / "player.js").read_text(encoding="utf-8")
    sections = []
    for i, svg in enumerate(pages, 1):
        body = "".join(svg).replace("url(#arrow)", f"url(#arrow{i})")
        sections.append(f'<section class="ssr-slide" aria-label="第 {i} 页"><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" role="img" aria-label="{escape(data["title"])} 第 {i} 页"><defs><marker id="arrow{i}" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="#{ACCENT}"/></marker></defs>{body}</svg><aside class="ssr-notes">{escape(data["notice"])}</aside></section>')
    html = f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(data["title"])}</title><style>{css}\n.ssr-slide svg{{display:block;width:100%;height:100%;font-family:"{FONT}",sans-serif}}\n</style></head><body><div class="ssr-player"><main class="ssr-stage"><div id="ssr-deck" data-width="1280" data-height="720" data-transition="none">{"".join(sections)}</div></main></div><script>{js}</script></body></html>'
    (out / "index.html").write_text(html, encoding="utf-8")
    (out / "edit-targets.json").write_text(json.dumps(targets, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "planning.md").write_text("# 双格式样例策划\n\n虚构数据用于维护验收；自主执行，4 页，HTML＋PPTX。\n\n视觉：暖纸底、深绿字、网格与适量留白；1280×720 对应 13⅓×7.5 英寸。使用本机 Microsoft YaHei，不分发系统字体。\n\n1. 上线进度与未确定事项：大数字、状态文字、独立概念插画。\n2. 门店补货机制：流程条件、退回补充和独立周报信息流。节点与文字可编辑、分组可移动；自动连线跟随未列为已验证能力。\n3. 方案比较：分组形状矩阵，成本与条件均保留；按格改字，不要求插行列。\n4. 实际上线趋势与待测收益：原生图表＋原生表格；预测不混入实际。\n\n来源：data.json 为虚构教学材料；illustration.png 由本样例绘制抽象图案，不是照片或实物证据。\n\n验收待执行：最终文件逐页预览、指定对象编辑、浏览器离线／打印。生成成功不代表通过；见运行后的 report.json 并补实际看图记录。\n", encoding="utf-8")
    print(f"Generated 4 slides and edit targets in {out}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    generate(parser.parse_args().out)
